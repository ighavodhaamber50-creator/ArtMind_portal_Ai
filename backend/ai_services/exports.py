"""Downloadable PDF and Word versions of a painting's details."""

import io
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm
from PIL import Image as PILImage
from reportlab.lib import colors as rl_colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

COLOR_HEX = {
    'blue': '#3342E8', 'green': '#1F9D6B', 'yellow': '#F5C518', 'red': '#D9362B',
    'orange': '#F28A1D', 'purple': '#7A4DD8', 'pink': '#F07BA6', 'white': '#F4F4F2',
    'brown': '#7A4E2D', 'teal': '#14A3A3', 'gold': '#D4AF37',
}
CATEGORY_HEX = {
    'abstract': '#3342E8', 'landscape': '#0E9F7E', 'flower': '#E8447D',
    'nature': '#6E9B1F', 'figurative': '#F5A400', 'religious': '#7A4DD8',
}
INK = '#1B1740'
SOFT = '#4A4670'
LINE = '#DCDEEA'


def _facts(artwork):
    return [
        ('Artist', artwork.artist.name),
        ('Category', artwork.category.name),
        ('Style', artwork.get_style_display()),
        ('Medium', artwork.get_medium_display()),
        ('Surface', artwork.get_surface_display()),
        ('Colors', ', '.join(artwork.dominant_colors) or 'Not recorded'),
    ]


def _summary_label(is_ai):
    return 'Summary (written by AI)' if is_ai else "Summary (built from the painting's details)"


def _fit(path, max_width, max_height):
    with PILImage.open(path) as image:
        width, height = image.size
    scale = min(max_width / width, max_height / height)
    return width * scale, height * scale


def build_pdf(artwork, summary, is_ai, image_path):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.6 * cm,
        bottomMargin=1.8 * cm, title=artwork.title, author='ArtMind AI Portal')

    base = getSampleStyleSheet()
    ink = rl_colors.HexColor(INK)
    title_style = ParagraphStyle('PaintingTitle', parent=base['Title'], fontName='Helvetica-Bold',
                                 fontSize=24, leading=28, textColor=ink, alignment=0, spaceBefore=12)
    byline = ParagraphStyle('Byline', parent=base['Normal'], fontName='Helvetica-Oblique',
                            fontSize=13, textColor=rl_colors.HexColor(SOFT), spaceAfter=10)
    heading = ParagraphStyle('Heading', parent=base['Heading2'], fontName='Helvetica-Bold',
                             fontSize=13, textColor=ink, spaceBefore=14, spaceAfter=4)
    body = ParagraphStyle('Body', parent=base['Normal'], fontName='Times-Roman', fontSize=11.5, leading=16)

    story = []

    # A bar in the category's colour across the top of the page
    bar = Table([['']], colWidths=[17 * cm], rowHeights=[0.35 * cm])
    bar.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1),
                              rl_colors.HexColor(CATEGORY_HEX.get(artwork.category.slug, INK)))]))
    story += [bar, Paragraph(escape(artwork.title), title_style),
              Paragraph(escape(artwork.artist.name), byline)]

    if image_path:
        width, height = _fit(image_path, 17 * cm, 11 * cm)
        story += [Image(image_path, width=width, height=height), Spacer(1, 10)]

    rows = [[Paragraph(f'<b>{escape(label)}</b>', body), Paragraph(escape(value), body)]
            for label, value in _facts(artwork)]
    facts = Table(rows, colWidths=[3.2 * cm, 13.8 * cm])
    facts.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LINEBELOW', (0, 0), (-1, -1), 0.25, rl_colors.HexColor(LINE)),
    ]))
    story.append(facts)

    if artwork.dominant_colors:
        swatches = Table([['' for _ in artwork.dominant_colors]],
                         colWidths=[1.4 * cm] * len(artwork.dominant_colors), rowHeights=[0.7 * cm])
        swatches.hAlign = 'LEFT'
        swatches.setStyle(TableStyle(
            [('BACKGROUND', (i, 0), (i, 0), rl_colors.HexColor(COLOR_HEX.get(name, '#8A8FA3')))
             for i, name in enumerate(artwork.dominant_colors)]
            + [('BOX', (0, 0), (-1, -1), 0.25, rl_colors.HexColor(LINE))]))
        story += [Spacer(1, 8), swatches]

    story += [Paragraph('Description', heading), Paragraph(escape(artwork.description), body),
              Paragraph(_summary_label(is_ai), heading), Paragraph(escape(summary), body)]

    artist = artwork.artist
    if artist.bio or artist.nationality:
        story.append(Paragraph(f'About {escape(artist.name)}', heading))
        if artist.nationality:
            story.append(Paragraph(escape(artist.nationality), byline))
        if artist.bio:
            story.append(Paragraph(escape(artist.bio), body))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont('Helvetica', 8.5)
        canvas.setFillColor(rl_colors.HexColor(SOFT))
        canvas.drawString(2 * cm, 1 * cm, 'ArtMind AI Portal')
        canvas.drawRightString(A4[0] - 2 * cm, 1 * cm, f'Page {document.page}')
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def build_docx(artwork, summary, is_ai, image_path):
    doc = Document()
    doc.core_properties.title = artwork.title
    doc.core_properties.author = 'ArtMind AI Portal'

    doc.add_heading(artwork.title, level=0)
    byline = doc.add_paragraph(artwork.artist.name)
    byline.runs[0].italic = True

    if image_path:
        doc.add_picture(image_path, width=Cm(15))

    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for label, value in _facts(artwork):
        cells = table.add_row().cells
        cells[0].text = label
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = value

    doc.add_heading('Description', level=2)
    doc.add_paragraph(artwork.description)
    doc.add_heading(_summary_label(is_ai), level=2)
    doc.add_paragraph(summary)

    artist = artwork.artist
    if artist.bio or artist.nationality:
        doc.add_heading(f'About {artist.name}', level=2)
        if artist.nationality:
            doc.add_paragraph(artist.nationality).runs[0].italic = True
        if artist.bio:
            doc.add_paragraph(artist.bio)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()